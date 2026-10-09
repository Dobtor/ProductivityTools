"""RenderOutput — 輸出（HTML / DOCX）。

元素樹 → HTML / DOCX。

前端有 canvas-editor 的 getHTML()，伺服器端沒有，所以匯出鏈以
content_json 為權威、在這裡自己轉。涵蓋本模組實際會產生的元素類型；
未知類型退回「輸出其 value 的逸出文字」——最壞情況是掉格式，不是掉內容。

攤平（_flatten_content_json）在這一層而不是快照層：它是輸出的一部分，
而且標記藥丸在這裡要再丟一次（見 _MARKER_SOURCES）。
"""
import html as html_mod
import io
import json


from odoo.exceptions import UserError


class RenderOutput:

    def _drop_marker_at_flatten(self, element):
        """攤平時要丟掉的標記藥丸。

        pageScope 是唯一的例外：它要活到攤平這一步，才轉成一個不印字的
        pageScope 元素，由 _elements_to_html 變成段落的 class。
        跟著其他標記一起丟掉的話，「只在首頁」那一段會變成每一頁都印
        ——而且完全不報錯。

        它仍然是標記（isMarker=True）：快照的純量迴圈照樣跳過它，不求值、
        不印標籤文字。差別只在「什麼時候消失」。
        """
        if not self._is_marker_element(element):
            return False
        meta = self._element_field_meta(element) or {}
        return (meta.get('source') or '') != self._PAGE_SCOPE_SOURCE

    def _flatten_content_json(self, tree):
        """把藥丸攤平成純文字元素（就地改寫並回傳 tree）。

        決策四（只輸出值）：丟掉 label 樣式與 extension，網底屬編輯輔助，
        不進正式文件。**不求值**——值應已由快照凍結在 value 裡。
        """
        if not tree:
            return tree
        for elements in self._iter_element_lists(tree):
            # 標記藥丸整個丟掉（見 _MARKER_SOURCES 的說明），pageScope 例外
            elements[:] = [el for el in elements
                           if not self._drop_marker_at_flatten(el)]
            for idx, el in enumerate(elements):
                if not isinstance(el, dict) or el.get('type') != 'label':
                    continue
                # meta 要在丟掉 extension 之前讀——圖片與頁碼的攤平結果
                # 取決於它們的 source
                meta = self._element_field_meta(el) or {}
                src = (meta.get('source') or '').strip()
                plain = {k: v for k, v in el.items()
                         if k not in ('type', 'label', 'extension', 'labelId')}
                plain['value'] = el.get('value') or ''
                if src == self._IMAGE_SOURCE:
                    if (plain['value'] or '').startswith('data:'):
                        plain['type'] = 'image'
                        for key in ('width', 'height'):
                            if meta.get(key):
                                plain[key] = meta[key]
                    else:
                        # 沒有圖（欄位空的）→ 什麼都不印，不要留下 data URI 殘骸
                        plain['value'] = ''
                elif src == self._HTML_SOURCE:
                    # 區塊級：不能當行內塞進 <p> 裡，否則變成
                    # <p><p>…</p></p> 的非法嵌套
                    plain['type'] = 'htmlBlock'
                elif src == self._PAGE_SCOPE_SOURCE:
                    # 不印字，只把「這一段限哪些頁」帶到組 HTML 那一步
                    plain['type'] = 'pageScope'
                    scope = (meta.get('scope') or '').strip()
                    plain['scope'] = (
                        scope if scope in self._PAGE_SCOPES else '')
                    plain['value'] = ''
                elif src == self._PAGE_SOURCE:
                    plain['type'] = 'pageField'
                    plain['pageKind'] = (
                        'count' if (meta.get('part') or '') == 'count'
                        else 'number'
                    )
                    plain['value'] = ''
                elements[idx] = plain
        return tree

    # ─── 元素樹 → HTML（伺服器端匯出鏈用）────────────────────────────
    #
    # 前端有 canvas-editor 的 getHTML()，伺服器端沒有。匯出改以 content_json
    # 為權威之後，這支轉換是必要的。涵蓋本模組實際會產生的元素類型；
    # 未知類型退回「輸出其 value 的逸出文字」——最壞情況是掉格式，不是掉內容。

    _ROW_FLEX_ALIGN = {
        'left': 'left', 'center': 'center', 'right': 'right',
        'alignment': 'justify', 'justify': 'justify',
    }

    def _element_style(self, el):
        parts = []
        if el.get('bold'):
            parts.append('font-weight:bold')
        if el.get('italic'):
            parts.append('font-style:italic')
        decos = []
        if el.get('underline'):
            decos.append('underline')
        if el.get('strikeout'):
            decos.append('line-through')
        if decos:
            parts.append('text-decoration:%s' % ' '.join(decos))
        if el.get('color'):
            parts.append('color:%s' % el['color'])
        if el.get('highlight'):
            parts.append('background-color:%s' % el['highlight'])
        if el.get('size'):
            # canvas-editor 的 element.size 單位是 px——它組 canvas font 字串時是
            # `${size}px`（見 lib 的 getElementFont）。這裡若寫 pt，匯出的每段文字
            # 都會比畫面上大 33%（16px → 16pt = 21.3px），而且不會有任何錯誤訊息。
            parts.append('font-size:%spx' % el['size'])
        if el.get('font'):
            parts.append("font-family:'%s'" % str(el['font']).replace("'", ''))
        return ';'.join(parts)

    def _inline_element_html(self, el):
        """單一行內元素 → HTML 片段。"""
        etype = el.get('type') or 'text'
        if etype == 'htmlBlock':
            # 正常會在 _elements_to_html 被攔成區塊級；走到這裡表示它被放進了
            # 超連結之類的子串列，原樣輸出總比印出逸出標籤好
            return el.get('value') or ''
        if etype == 'pageField':
            # 頁碼不必自己算也不必寫 JS：Odoo 的 wkhtmltopdf 管線會把抽出來的
            # div.footer 包進 web.minimal_layout（subst=True，見
            # ir_actions_report.py:441），那支範本內的 subst() JS 會填滿所有
            # class="page" / "topage" 的元素。輸出這個 span 就夠了。
            kind = el.get('pageKind') or 'number'
            return '<span class="%s"></span>' % (
                'topage' if kind == 'count' else 'page'
            )
        if etype == 'image':
            src = el.get('value') or ''
            if not src:
                return ''
            w = el.get('width')
            h = el.get('height')
            dims = ''
            if w:
                dims += ' width="%d"' % int(w)
            if h:
                dims += ' height="%d"' % int(h)
            return '<img src="%s"%s/>' % (html_mod.escape(src, quote=True), dims)
        if etype == 'separator':
            return '<hr/>'
        if etype == 'tab':
            return '&emsp;'
        if etype == 'checkbox':
            checked = ((el.get('checkbox') or {}).get('value'))
            return '☑' if checked else '☐'
        if etype == 'radio':
            checked = ((el.get('radio') or {}).get('value'))
            return '◉' if checked else '○'
        if etype == 'hyperlink':
            inner = ''.join(
                self._inline_element_html(c)
                for c in (el.get('valueList') or []) if isinstance(c, dict)
            ) or html_mod.escape(el.get('value') or '')
            url = html_mod.escape(el.get('url') or '', quote=True)
            return '<a href="%s">%s</a>' % (url, inner)

        text = html_mod.escape(el.get('value') or '')
        if '\n' in text:
            # 伺服器端求值會產生多行值（format_address、多行文字欄位）。
            # canvas-editor 本身不把換行放進元素的 value 裡（它用獨立的 '\n'
            # 元素），而「值剛好等於 '\n'」的段落標記在 _elements_to_html 就先
            # 攔掉了——所以走到這裡的換行一定是值內部的，轉成 <br/> 沒有歧義。
            # 不轉的話多行地址會塌成一行，而 HTML 裡看起來只是少了換行。
            text = text.replace('\n', '<br/>')
        if not text:
            return ''
        style = self._element_style(el)
        if etype == 'label':
            # 理論上匯出前已攤平；萬一漏了也只輸出文字，不帶網底（決策四）
            return text
        return '<span style="%s">%s</span>' % (style, text) if style else text

    def _table_to_html(self, el):
        rows = []
        for row in (el.get('trList') or []):
            if not isinstance(row, dict):
                continue
            cells = []
            for cell in (row.get('tdList') or []):
                if not isinstance(cell, dict):
                    continue
                attrs = ''
                if (cell.get('colspan') or 1) > 1:
                    attrs += ' colspan="%d"' % int(cell['colspan'])
                if (cell.get('rowspan') or 1) > 1:
                    attrs += ' rowspan="%d"' % int(cell['rowspan'])
                cells.append('<td%s>%s</td>' % (
                    attrs, self._elements_to_html(cell.get('value') or []),
                ))
            if cells:
                rows.append('<tr>%s</tr>' % ''.join(cells))
        if not rows:
            return ''
        # 區塊容器不是真的表格，不可印框線。class 讓四條輸出路徑各自關掉邊框，
        # 而不是在這裡寫 inline style——inline style 進不了 python-docx 的表格樣式。
        cls = ' class="doc-block"' if self._element_block_kind(el) else ''
        return '<table%s>%s</table>' % (cls, ''.join(rows))

    def _elements_to_html(self, elements):
        """元素串列 → HTML。

        canvas-editor 的元素串列是扁平的，用 value == '\\n' 標示換行；
        換行元素身上的 rowFlex 決定該段落的對齊。
        """
        out = []
        buf = []
        # 這一段的頁面範圍（由段落內的 pageScope 元素設定）。每沖一次就清掉。
        scope = {'value': ''}

        def _flush(row_el=None):
            cls = (' class="doc-page-%s"' % scope['value']) if scope['value'] else ''
            scope['value'] = ''
            if not buf:
                # 空段落也要保留，否則多個空行會被吃掉、版面走樣
                out.append('<p%s><br/></p>' % cls)
                return
            align = self._ROW_FLEX_ALIGN.get((row_el or {}).get('rowFlex') or '')
            style = ' style="text-align:%s"' % align if align else ''
            out.append('<p%s%s>%s</p>' % (cls, style, ''.join(buf)))
            buf.clear()

        for el in (elements or []):
            if not isinstance(el, dict):
                continue
            etype = el.get('type') or 'text'
            if etype == 'table':
                if buf:
                    _flush()
                out.append(self._table_to_html(el))
                continue
            if etype == 'htmlBlock':
                # 已在求值時消毒過（_html_field_value）。這裡原樣輸出，
                # 與表格同樣是區塊級：先把行內緩衝沖掉再放。
                if buf:
                    _flush()
                out.append(el.get('value') or '')
                continue
            if etype == 'pageScope':
                # 標記本身不印字；它只決定所在段落的 class
                if el.get('scope'):
                    scope['value'] = el['scope']
                continue
            if etype == 'pageBreak':
                if buf:
                    _flush()
                out.append('<div class="doc-page-break"></div>')
                continue
            if (el.get('value') or '') == '\n':
                _flush(el)
                continue
            buf.append(self._inline_element_html(el))
        if buf:
            _flush()
        return ''.join(out)

    # ─── HTML → content_json（_content_json_to_html 的反向）──────────
    #
    # 為什麼需要：模組自己出貨的 6 張 data 範本只有 content_html，沒有
    # content_json。它們列印時走 _render_template 那條舊路，於是**享受不到任何
    # 藥丸功能**——型別格式、欄位標籤、頁面範圍、條件、重複列都不生效。
    #
    # ☠️ Phase 5 的註解寫著「HTML → IElement 需要 canvas-editor 的
    # executeSetHTML，那是瀏覽器端的東西」——對**任意** HTML 是對的。
    # 這支刻意只處理**我們自己寫的那個子集**：出貨範本實際用到的標籤只有
    #   p / h1 / h2 / h3 / b / strong / br / table / thead / tbody / tr / th / td
    #（數過的）。遇到子集外的標籤會留下 note，呼叫端決定要不要接受。
    #
    # 正確性靠**來回轉換**釘住：html → content_json → html 的結果要與原 html
    # 等價（見 TestHtmlToContentJson 的 round-trip 測試）。那比逐個標籤寫斷言
    # 可靠，因為反向那條路（_content_json_to_html）是已經在用的。

    _HTML_BLOCK_TAGS = ('p', 'h1', 'h2', 'h3', 'div')
    # 標題的字級（px）。canvas-editor 的 element.size 單位是 px，
    # 與 _element_style 的 font-size:%spx 對齊。
    _HTML_HEADING_SIZE = {'h1': 24, 'h2': 20, 'h3': 16}
    _HTML_ALIGN_FLEX = {'center': 'center', 'right': 'right',
                        'justify': 'alignment', 'left': 'left'}

    def _html_to_content_json(self, html, notes=None):
        """把一段（我們自己寫的）HTML 轉成 content_json 的 main 元素串列。

        回傳 {'header': [], 'main': [...], 'footer': []}。
        notes 給一個 list 的話，遇到不支援的標籤會 append 說明而不是靜默吞掉。
        """
        from lxml import html as lhtml
        notes = notes if notes is not None else []
        text = (html or '').strip()
        if not text:
            return {'header': [], 'main': [], 'footer': []}
        root = lhtml.fragment_fromstring(text, create_parent='div')
        main = []
        self._html_walk_blocks(root, main, notes)
        return {'header': [], 'main': main, 'footer': []}

    def _html_walk_blocks(self, parent, out, notes):
        """走訪區塊層：每個區塊產生它的行內元素，再補一個 '\n'。"""
        for node in parent:
            if not isinstance(node.tag, str):
                # ☠️ 註解（與 PI）的 .tag 不是字串而是一個 callable。不擋的話會
                # 掉到下面的「未知標籤」分支，把**註解內文當成正文輸出**——
                # 出貨範本裡就有一段 <!-- Sprint Y12.2… --> 的註解，實測結果是
                # 那張範本印出註解文字、而且後面的內容整段不見。
                # 註解自己丟掉，它後面的文字（tail）要留。
                if node.tail and node.tail.strip():
                    out.append({'value': node.tail})
                continue
            tag = (node.tag or '').lower()
            if tag == 'table':
                out.append(self._html_table_to_element(node, notes))
                continue
            if tag in self._HTML_BLOCK_TAGS:
                runs = []
                self._html_walk_inline(node, runs, {}, notes)
                size = self._HTML_HEADING_SIZE.get(tag)
                if size:
                    for r in runs:
                        r.setdefault('size', size)
                        r.setdefault('bold', True)
                out.extend(runs)
                out.append(self._html_newline(node))
                # div 可能自己包著區塊（出貨範本沒有，但容錯）
                if tag == 'div':
                    self._html_walk_blocks(node, out, notes)
                continue
            if tag == 'br':
                out.append({'value': '\n'})
                continue
            if tag:
                notes.append('不支援的標籤 <%s>，內容以純文字保留' % tag)
            runs = []
            self._html_walk_inline(node, runs, {}, notes)
            out.extend(runs)

    def _html_newline(self, node):
        """區塊結尾的 '\n' 元素；對齊寫在它身上（與 _elements_to_html 相反方向）。"""
        nl = {'value': '\n'}
        style = (node.get('style') or '')
        for css, flex in self._HTML_ALIGN_FLEX.items():
            if 'text-align:%s' % css in style.replace(' ', ''):
                nl['rowFlex'] = flex
                break
        return nl

    def _html_walk_inline(self, node, out, inherited, notes):
        """走訪行內層，把 b / strong 等轉成元素身上的屬性。"""
        if node.text:
            out.append(dict(inherited, value=node.text))
        for child in node:
            if not isinstance(child.tag, str):
                # 同上：註解不是內容
                if child.tail:
                    out.append(dict(inherited, value=child.tail))
                continue
            tag = (child.tag or '').lower()
            attrs = dict(inherited)
            if tag in ('b', 'strong'):
                attrs['bold'] = True
            elif tag in ('i', 'em'):
                attrs['italic'] = True
            elif tag == 'u':
                attrs['underline'] = True
            elif tag == 'br':
                out.append({'value': '\n'})
                if child.tail:
                    out.append(dict(inherited, value=child.tail))
                continue
            elif tag and tag not in ('span', 'font'):
                notes.append('不支援的行內標籤 <%s>，內容以純文字保留' % tag)
            self._html_walk_inline(child, out, attrs, notes)
            if child.tail:
                out.append(dict(inherited, value=child.tail))

    def _html_table_to_element(self, table, notes):
        """<table> → {type:'table', trList, colgroup}。"""
        rows = []
        for tr in table.iter('tr'):
            cells = []
            for td in tr:
                if not isinstance(td.tag, str):
                    continue
                tag = (td.tag or '').lower()
                if tag not in ('td', 'th'):
                    continue
                runs = []
                # 儲存格內可以有區塊（<p>）也可以直接是文字
                if any(isinstance(c.tag, str)
                       and c.tag.lower() in self._HTML_BLOCK_TAGS for c in td):
                    self._html_walk_blocks(td, runs, notes)
                else:
                    self._html_walk_inline(td, runs, {}, notes)
                    runs.append({'value': '\n'})
                if tag == 'th':
                    for r in runs:
                        if r.get('value') != '\n':
                            r['bold'] = True
                cell = {'value': runs}
                for attr, key in (('colspan', 'colspan'), ('rowspan', 'rowspan')):
                    try:
                        v = int(td.get(attr) or 1)
                    except ValueError:
                        v = 1
                    if v > 1:
                        cell[key] = v
                cells.append(cell)
            if cells:
                rows.append({'tdList': cells})
        width = max((len(r['tdList']) for r in rows), default=0)
        el = {'type': 'table', 'value': '', 'trList': rows}
        if width:
            # 平均分配欄寬。canvas-editor 需要 colgroup 才畫得出表格，
            # 而原 HTML 的 width 是百分比、不是它要的 px。
            el['colgroup'] = [{'width': int(718 / width)} for _ in range(width)]
        return el

    def _content_json_to_html(self, tree, zone='main'):
        """content_json 的指定區域 → HTML。"""
        if isinstance(tree, dict):
            elements = tree.get(zone)
        elif isinstance(tree, list) and zone == 'main':
            elements = tree
        else:
            elements = None
        return self._elements_to_html(elements or [])

    def _parse_content_json(self, raw):
        """content_json 欄位（字串或已解析物件）→ dict/list；失敗回 None。"""
        if not raw:
            return None
        if isinstance(raw, (dict, list)):
            return raw
        try:
            return json.loads(raw)
        except Exception:
            return None

    # ─── HTML → DOCX（doc.document 與 doc.output 共用）────────────────
    #
    # 原本只存在 doc.document._generate_docx_via_python 裡，綁在它的欄位上。
    # 報表引擎的輸出紀錄也要能匯出 DOCX（定案決策四：定位為「匯出去編輯」
    # 的便利功能），所以抽到 mixin、參數化。
    # doc.document 那支保留為對外介面，內部委派到這裡。

    _DOCX_PAGE_SIZES_MM = {
        'A4': (210, 297), 'A3': (297, 420), 'A5': (148, 210),
        'letter': (216, 279), 'legal': (216, 356),
    }

    def _docx_fill_zone(self, zone, content, lhtml, add_runs, apply_align):
        """把頁首／頁尾的 HTML（或純文字）填進 DOCX 的 header/footer。"""
        if not content or not content.strip():
            return
        text = content.strip()
        if '<' not in text:
            zone.paragraphs[0].text = text
            return
        try:
            node = lhtml.fromstring('<div>%s</div>' % text)
        except Exception:
            zone.paragraphs[0].text = text
            return
        blocks = node.xpath('./p | ./div') or [node]
        for idx, blk in enumerate(blocks):
            para = zone.paragraphs[0] if idx == 0 else zone.add_paragraph()
            apply_align(para, blk)
            add_runs(para, blk)

    # 下載檔名不能有的字。'/' 是真的會壞（單號 S00001/2026 很常見），
    # 其餘是 Windows 的保留字元。
    _FILENAME_BAD_CHARS = '/\\:*?"<>|\r\n\t'

    def _render_filename(self, pattern, record):
        """把檔名樣式算成一個可以當檔名的字串；算不出來回空字串。

        用同一套沙箱（所以樣式寫 {{ object.name }} 與範本裡一致），而不是
        Odoo 的 safe_eval——使用者學一種語法就好。

        回空字串而不是拋例外：呼叫端會退回原生檔名。為了一個檔名讓下載失敗
        是最糟的結果。
        """
        pattern = (pattern or '').strip()
        if not pattern or record is None:
            return ''
        try:
            env_j = self._get_sandbox_env(record)
            text = env_j.from_string(pattern).render(
                object=record, user=self.env.user)
        except Exception:
            return ''
        text = (text or '').strip()
        if text in ('', 'False', 'None'):
            return ''
        for ch in self._FILENAME_BAD_CHARS:
            text = text.replace(ch, '_')
        # 連續底線收成一個，首尾的去掉——'/' 被取代後常常留下一串
        while '__' in text:
            text = text.replace('__', '_')
        text = text.strip('_ ')
        # 檔案系統的上限是 255 bytes，中文一個字 3 bytes；留副檔名的空間
        return text[:80] or ''

    def _docx_bytes_from_html(self, body_html, page_format='A4', margins=None,
                              header_text='', footer_text=''):
        """body HTML → DOCX bytes。

        margins：{'top','bottom','left','right'} 單位 px（與編輯器一致）。
        缺 python-docx 時 raise UserError 並指名該裝的 pip 套件——
        它是選用相依（見 __manifest__.py），核心功能不需要。
        """
        try:
            from docx import Document
            from docx.shared import Mm
            from lxml import html as lhtml
        except ImportError as e:
            raise UserError(
                f'無法匯出 DOCX：{e}\n'
                '請安裝 python-docx：pip install python-docx'
            )
        from ..doc_document import (
            _add_runs_from_node, _apply_paragraph_align, _html_node_to_docx,
        )

        margins = margins or {}
        px_to_mm = 0.264583  # 96dpi：1px = 0.264583mm
        w_mm, h_mm = self._DOCX_PAGE_SIZES_MM.get(page_format, (210, 297))

        docx_doc = Document()
        section = docx_doc.sections[0]
        section.page_width = Mm(w_mm)
        section.page_height = Mm(h_mm)
        section.top_margin = Mm(margins.get('top', 96) * px_to_mm)
        section.bottom_margin = Mm(margins.get('bottom', 96) * px_to_mm)
        section.left_margin = Mm(margins.get('left', 96) * px_to_mm)
        section.right_margin = Mm(margins.get('right', 96) * px_to_mm)

        try:
            tree = lhtml.fromstring('<div>%s</div>' % (body_html or ''))
            _html_node_to_docx(docx_doc, tree)
        except Exception:
            # 解析失敗時退回純文字——掉格式比整份匯不出來好
            try:
                text = lhtml.fromstring(
                    '<div>%s</div>' % (body_html or '')
                ).text_content()
            except Exception:
                text = ''
            docx_doc.add_paragraph(text)

        # 頁首頁尾收 HTML 而不是純文字：頁碼在 DOCX 裡是 field code
        # （<span class="page"> → PAGE），用 paragraphs[0].text 設值會把它
        # 變成字面文字「頁碼」，而且不會有任何錯誤訊息。
        self._docx_fill_zone(section.header, header_text, lhtml,
                             _add_runs_from_node, _apply_paragraph_align)
        self._docx_fill_zone(section.footer, footer_text, lhtml,
                             _add_runs_from_node, _apply_paragraph_align)

        buf = io.BytesIO()
        docx_doc.save(buf)
        return buf.getvalue()
