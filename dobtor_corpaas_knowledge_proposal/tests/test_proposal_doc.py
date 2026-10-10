# -*- coding: utf-8 -*-
"""服務建議書的文件輸出（dobtor_doc_editor）。

本機沒有 wkhtmltopdf／LibreOffice，所以 PDF 不測；HTML 路徑與 DOCX 轉換（python-docx）
都實際跑。
"""
import base64
import io
import re

from odoo.tests import TransactionCase, tagged

from ..models.proposal_doc import BINDING_XMLID, REPORT_XMLID, TEMPLATE_XMLID
from ..services import doc_template, docx_polish


def _walk(obj):
    """走過範本裡每一個 dict。"""
    if isinstance(obj, dict):
        yield obj
        for value in obj.values():
            yield from _walk(value)
    elif isinstance(obj, list):
        for item in obj:
            yield from _walk(item)


@tagged('post_install', '-at_install')
class TestDocTemplatePure(TransactionCase):

    def setUp(self):
        super().setUp()
        self.tree = doc_template.build()

    def _metas(self):
        return [el['extension']['dobtorField'] for el in _walk(self.tree)
                if isinstance(el.get('extension'), dict) and 'dobtorField' in el['extension']]

    def test_structure_is_json_serializable_and_has_three_zones(self):
        self.assertEqual(set(self.tree), {'header', 'main', 'footer'})
        import json
        self.assertEqual(json.loads(doc_template.build_json()), self.tree)

    def test_every_table_fits_the_page_and_is_rectangular(self):
        tables = [el for el in _walk(self.tree) if el.get('type') == 'table']
        self.assertGreaterEqual(len(tables), 12)
        for table in tables:
            widths = [c['width'] for c in table['colgroup']]
            self.assertLessEqual(sum(widths), doc_template.WIDTH, widths)
            for row in table['trList']:
                self.assertEqual(len(row['tdList']), len(widths))

    def test_every_table_is_followed_by_a_paragraph_break(self):
        """☠️ 表格後面緊接條件段落會被當成同一個段落：條件為假，前面那張表一起消失。"""
        main = self.tree['main']
        for idx, el in enumerate(main):
            if el.get('type') == 'table':
                self.assertLess(idx + 1, len(main))
                self.assertEqual(main[idx + 1].get('value'), '\n',
                                 '第 %d 個元素的表格後面不是換行' % idx)

    def test_text_elements_carry_the_house_font_and_body_size(self):
        """字型寫在每個元素上，doc_editor 轉 DOCX 才會連東亞字型一起設。"""
        elements = [el for el in _walk(self.tree)
                    if isinstance(el.get('value'), str) and el['value'] not in ('', '\n')
                    and el.get('type') not in ('table', 'pageBreak')]
        self.assertTrue(elements)
        for el in elements:
            self.assertEqual(el.get('font'), doc_template.FONT, el)
            self.assertTrue(el.get('size'), el)

    def test_cover_is_centered_and_followed_by_a_page_break(self):
        main = self.tree['main']
        breaks = [i for i, el in enumerate(main) if el.get('type') == 'pageBreak']
        self.assertEqual(len(breaks), 1, '封面後分頁一次')
        cover = main[:breaks[0]]
        centered = [el for el in cover if el.get('value') == '\n' and el.get('rowFlex') == 'center']
        self.assertGreaterEqual(len(centered), 6)

    def test_footer_has_company_version_and_a_page_number(self):
        metas = [(el.get('extension') or {}).get('dobtorField', {}) for el in self.tree['footer']]
        sources = [m.get('source') for m in metas if m]
        self.assertIn('page', sources, '頁尾要有頁碼藥丸')
        self.assertIn('expression', sources)

    def test_every_expression_reads_data_not_object(self):
        """範本只讀 data.*：已送出版本才能完全由凍結快照決定輸出。"""
        for meta in self._metas():
            for key in ('expression', 'sourceExpression'):
                expr = meta.get(key)
                if expr:
                    self.assertIn('data.', expr, '%s 沒有讀 data.*' % expr)
                    self.assertNotIn('object.', expr)

    def test_every_chapter_heading_and_header_row_is_conditional(self):
        """沒資料的章節要整章消失：標題段落與表頭列都要有條件標記。"""
        tables = [el for el in _walk(self.tree) if el.get('type') == 'table']
        for table in tables:
            if len(table['colgroup']) == 2 and len(table['trList']) > 2:
                continue            # 鍵值表
            header = table['trList'][0]
            sources = [(el.get('extension') or {}).get('dobtorField', {}).get('source')
                       for cell in header['tdList'] for el in cell['value']]
            self.assertIn('condition', sources, '表頭列少了條件標記：%s' % sources)

    def test_each_repeat_row_declares_a_source_and_lines_have_paths(self):
        repeats = [m for m in self._metas() if m['source'] == 'repeat']
        self.assertGreaterEqual(len(repeats), 12)
        self.assertEqual(len({m['repeatId'] for m in repeats}), len(repeats))
        for meta in repeats:
            self.assertTrue(meta['sourceExpression'].startswith('data.'))
        for meta in (m for m in self._metas() if m['source'] == 'line'):
            self.assertTrue(meta.get('path'))


@tagged('post_install', '-at_install')
class TestProposalDoc(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        icp = cls.env['ir.config_parameter'].sudo()
        for key, value in {'day_rate': 10000, 'internal_day_cost': 4000, 'hours_per_day': 8,
                           'hours_per_week': 20, 'validity_days': 60, 'warranty_months': 3,
                           'payment_days': 10, 'mobilization_ratio': 30,
                           'acceptance_days': 10}.items():
            icp.set_param('corpaas_proposal.%s' % key, value)
        cls.Proposal = cls.env['corpaas.knowledge.proposal']
        # force：這個資料庫可能是舊版範本升級來的；測試要驗的是現在出貨的版面
        cls.Proposal._ensure_doc_template(force=True)
        cls.partner = cls.env['res.partner'].create({'name': '文件客戶股份有限公司'})
        Cap = cls.env['corpaas.knowledge.capability']
        cls.cap_a = Cap.create({'name': '財務會計', 'code': 'doc_fin', 'color': 'native',
                                'pain': '月結對帳要三天', 'outcome': '自動對帳，一天結帳'})
        cls.cap_b = Cap.create({'name': '人資差勤', 'code': 'doc_hr', 'color': 'dobtor',
                                'pain': '請假靠紙本', 'outcome': '線上請假與假勤彙總'})
        Tpl = cls.env['corpaas.knowledge.effort_template']
        Tpl.create({'capability_id': cls.cap_a.id, 'activity': 'requirements', 'base_days': 3})
        Tpl.create({'capability_id': cls.cap_b.id, 'activity': 'requirements', 'base_days': 2})
        Tpl.create({'activity': 'pm', 'base_days': 2})

    def _full(self, **vals):
        """時間制、全部組成都有的版本。"""
        prop = self.Proposal.create(dict({
            'partner_id': self.partner.id, 'pricing_model': 'time_budget',
            'hourly_rate': 2500, 'budget_cap': 300000, 'company_count': 2, 'user_count': 20,
            'doc_purpose': '導入範圍確認與預算配比'}, **vals))
        Pain, Mapping = self.env['corpaas.knowledge.pain'], self.env['corpaas.knowledge.mapping']
        p1 = Pain.create({'proposal_id': prop.id, 'department': '財務', 'description': '月結要三天'})
        p2 = Pain.create({'proposal_id': prop.id, 'department': '人資', 'description': '請假靠紙本'})
        p3 = Pain.create({'proposal_id': prop.id, 'description': '要客製報表'})
        Mapping.create([
            {'pain_id': p1.id, 'capability_id': self.cap_a.id, 'color': 'native',
             'confirmed': True, 'note': '自動對帳'},
            {'pain_id': p2.id, 'capability_id': self.cap_b.id, 'color': 'dobtor',
             'confirmed': True},
            {'pain_id': p3.id, 'color': 'custom', 'custom_days_low': 3, 'custom_days_high': 5,
             'confirmed': True}])
        prop.action_compute_estimate()
        prop.action_suggest_units()
        prop.action_suggest_payments()
        prop.action_suggest_phases()
        prop.action_load_clauses()
        self.env['corpaas.knowledge.basis'].create({
            'proposal_id': prop.id, 'kind': 'file', 'title': 'Requirements 匯出檔',
            'date': '2026-08-31'})
        self.env['corpaas.knowledge.change_line'].create({
            'proposal_id': prop.id, 'ref_no': '1', 'demand': '財務範圍限應收付',
            'handling': '本版限縮財務範圍', 'chapter_ref': '第三章'})
        self.env['corpaas.knowledge.scope_item'].create([
            {'proposal_id': prop.id, 'kind': 'in', 'name': '應收付管理'},
            {'proposal_id': prop.id, 'kind': 'out', 'name': '總帳作業'}])
        self.env['corpaas.knowledge.obligation'].create({
            'proposal_id': prop.id, 'kind': 'client', 'text': '提供科目表'})
        self.env['corpaas.knowledge.benefit_item'].create({
            'proposal_id': prop.id, 'name': '月結對帳', 'current_manual': '3 天',
            'after': '1 天', 'annual_saving': 120000})
        return prop

    def _html(self, prop):
        html, kind = self.env['ir.actions.report']._render_qweb_html(REPORT_XMLID, prop.ids)
        self.assertEqual(kind, 'html')
        return html if isinstance(html, str) else html.decode()

    # ------------------------------------------------------------------
    def test_template_and_binding_are_installed_once(self):
        tmpl = self.env.ref(TEMPLATE_XMLID)
        binding = self.env.ref(BINDING_XMLID)
        self.assertEqual(tmpl.model_id.model, 'corpaas.knowledge.proposal')
        self.assertEqual(binding.template_id, tmpl)
        self.assertEqual(binding.report_id, self.env.ref(REPORT_XMLID))
        self.assertTrue(binding.persist_output, '寄給客戶的文件每次輸出都要留紀錄')
        self.assertEqual(self.Proposal._ensure_doc_template(), tmpl, '冪等')
        self.assertEqual(self.env['doc.report'].sudo().search_count(
            [('report_id', '=', binding.report_id.id)]), 1)

    def test_force_rebuilds_the_shipped_layout_only_when_asked(self):
        tmpl = self.env.ref(TEMPLATE_XMLID)
        tmpl.sudo().content_json = '{"header":[],"main":[],"footer":[]}'
        self.Proposal._ensure_doc_template()
        self.assertEqual(tmpl.content_json, '{"header":[],"main":[],"footer":[]}',
                         '客戶改過的版面升級時不能被蓋掉')
        self.Proposal._ensure_doc_template(force=True)
        self.assertGreater(len(tmpl.content_json), 5000)

    def test_data_numbers_only_the_chapters_that_have_content(self):
        prop = self._full()
        data = prop.doc_report_values()
        chap = data['chap']
        self.assertEqual(list(chap.values()), ['一', '二', '三', '四', '五', '六', '七',
                                               '八', '九', '十', '十一', '十二'][:len(chap)],
                         '編號必須連續，沒資料的章節不佔號')
        self.assertIn('budget', chap)
        self.assertIn('clauses', chap)
        # 沒有任何選用組成的版本：只剩有內容的章節
        bare = self.Proposal.create({'partner_id': self.partner.id})
        self.assertNotIn('changes', bare.doc_report_values()['chap'])
        self.assertNotIn('benefits', bare.doc_report_values()['chap'])

    def test_every_data_key_the_template_reads_exists(self):
        """範本與 doc_report_values() 的契約：範本讀的鍵，資料都要有。"""
        data = self._full().doc_report_values()
        tree = doc_template.build()
        paths = set()
        for el in _walk(tree):
            meta = (el.get('extension') or {}).get('dobtorField') or {}
            for key in ('expression', 'sourceExpression'):
                for m in re.finditer(r'data\.(\w+)(?:\.(\w+))?', meta.get(key) or ''):
                    paths.add((m.group(1), m.group(2)))
        self.assertTrue(paths)
        for top, sub in sorted(paths, key=str):
            self.assertIn(top, data, 'data.%s 不存在' % top)
            # chap 只有「有內容的章節」才有編號，範本用 `or ''` 容許缺鍵
            if sub and isinstance(data[top], dict) and top != 'chap':
                self.assertIn(sub, data[top], 'data.%s.%s 不存在' % (top, sub))

    def test_repeat_line_paths_exist_on_every_row(self):
        data = self._full().doc_report_values()
        tree = doc_template.build()
        by_id = {}
        for el in _walk(tree):
            meta = (el.get('extension') or {}).get('dobtorField') or {}
            if meta.get('source') == 'repeat':
                by_id[meta['repeatId']] = meta['sourceExpression'].split('.', 1)[1]
        for table in (el for el in _walk(tree) if el.get('type') == 'table'):
            if len(table['trList']) != 2:
                continue
            body = table['trList'][1]
            rid = next(((e['extension']['dobtorField']['repeatId']) for c in body['tdList']
                        for e in c['value'] if (e.get('extension') or {}).get(
                            'dobtorField', {}).get('source') == 'repeat'), None)
            if not rid:
                continue
            rows = data[by_id[rid]]
            fields = [e['extension']['dobtorField']['path'] for c in body['tdList']
                      for e in c['value'] if (e.get('extension') or {}).get(
                          'dobtorField', {}).get('source') == 'line']
            for row in rows:
                for f in fields:
                    self.assertIn(f, row, '%s 的列少了 %s' % (by_id[rid], f))

    # ------------------------------------------------------------------
    def test_html_has_the_cover_chapters_and_content(self):
        prop = self._full()
        html = self._html(prop)
        for needle in ('文件客戶股份有限公司', '服務建議書', 'v1.0', '導入範圍確認與預算配比',
                       '2026-08-31 Requirements 匯出檔', '一、需求核對與本版處理',
                       '財務範圍限應收付', '月結要三天', '自動對帳', '應收付管理', '總帳作業',
                       '預算規劃', '動員款', '驗收簽認', '備註與免責', '效益評估',
                       '提供科目表', '120,000'):
            self.assertIn(needle, html)
        # 標記藥丸是宣告，不是內容；設定錯誤才會原樣留著
        for leak in ('條件', '明細'):
            # 標記藥丸的標籤單獨成詞；「前提條件」「計費明細」這類合法文字前後都有字
            self.assertIsNone(re.search(r'(?<![\u4e00-\u9fff])%s(?![\u4e00-\u9fff])' % leak, html),
                              '輸出裡漏出了範本標記：%s' % leak)
        for leak in ('{{', 'data.'):
            self.assertNotIn(leak, html, '輸出裡漏出了範本標記：%s' % leak)

    def test_chapters_without_data_disappear_and_numbers_stay_consecutive(self):
        prop = self.Proposal.create({
            'partner_id': self.partner.id, 'pricing_model': 'time_budget', 'hourly_rate': 2500})
        pain = self.env['corpaas.knowledge.pain'].create(
            {'proposal_id': prop.id, 'description': '月結要三天'})
        self.env['corpaas.knowledge.mapping'].create({
            'pain_id': pain.id, 'capability_id': self.cap_a.id, 'color': 'native',
            'confirmed': True})
        prop.action_compute_estimate()
        html = self._html(prop)
        self.assertIn('月結要三天', html)
        for gone in ('需求核對與本版處理', '範圍界線', '效益評估', '備註與免責',
                     '驗收與付款', '導入期程與里程碑'):
            self.assertNotIn(gone, html)
        chap = prop.doc_report_values()['chap']
        self.assertTrue(html.count('、') >= len(chap))
        numerals = [chap[k] + '、' for k in chap]
        positions = [html.index(n) for n in numerals]
        self.assertEqual(positions, sorted(positions), '章節編號依序出現')

    def test_a_table_survives_a_false_conditional_paragraph_right_after_it(self):
        """有驗收單元、沒有款別：驗收單元表不能被後面那個為假的「付款款別」小標題帶走。"""
        prop = self._full()
        prop.payment_ids.unlink()
        html = self._html(prop)
        self.assertIn('驗收單元與標準', html)
        self.assertIn('驗收標準（須全數通過）', html)
        self.assertNotIn('付款款別', html)
        # 工時制的預算大類表（後面緊接著為假的「計費明細」）
        full = self._html(self._full())
        self.assertIn('預算大類', full)
        self.assertNotIn('計費明細', full)

    def test_subscription_model_prints_the_billing_table_not_the_budget_split(self):
        prop = self.Proposal.create({'partner_id': self.partner.id})
        pain = self.env['corpaas.knowledge.pain'].create(
            {'proposal_id': prop.id, 'description': '要一個系統'})
        self.env['corpaas.knowledge.mapping'].create({
            'pain_id': pain.id, 'capability_id': self.cap_a.id, 'color': 'native',
            'confirmed': True})
        prop.action_compute_estimate()
        data = prop.doc_report_values()
        self.assertTrue(data['show']['budget_sub'])
        self.assertFalse(data['show']['budget_time'])
        html = self._html(prop)
        self.assertIn('計費明細', html)
        self.assertNotIn('工時與預算配比', html)

    def test_sent_version_prints_the_frozen_snapshot(self):
        """送出後不論能力文案怎麼改，重新輸出的永遠是送出時的那份。"""
        prop = self._full()
        prop.action_send()
        before = self._html(prop)
        self.assertIn('月結對帳要三天', before)
        self.cap_a.pain = '（送出後才改的文案）'
        self.cap_b.name = '送出後改名'
        after = self._html(prop)
        self.assertEqual(before, after, '同一版重新輸出必須一致')
        self.assertNotIn('送出後才改的文案', after)
        # 草稿則讀現場資料
        draft = self._full()
        self.assertIn('送出後才改的文案', self._html(draft))

    def test_new_version_prints_its_own_number_and_change_log(self):
        v10 = self._full()
        v10.action_send()
        v11 = self.Proposal.browse(v10.action_copy_new_version()['res_id'])
        v11.action_draft_changes()
        html = self._html(v11)
        self.assertIn('v1.1', html)
        self.assertNotIn('財務範圍限應收付', html, '前一版的核對項不繼承')
        self.assertEqual(v11.doc_report_values()['cover']['version_no'], 'v1.1')

    def test_old_snapshots_without_the_new_sections_still_render(self):
        """1.x 送出的版本：快照沒有新的組成，輸出不能壞。"""
        prop = self._full()
        prop.action_send()
        import json
        snap = json.loads(prop.snapshot_json)
        for key in ('version', 'pricing', 'basis', 'changes', 'scope', 'obligations',
                    'phases', 'units', 'payments', 'clauses', 'benefits', 'case', 'company'):
            snap.pop(key, None)
        prop._internal().write({'snapshot_json': json.dumps(snap)})
        data = prop.doc_report_values()
        self.assertFalse(data['show']['changes'])
        self.assertFalse(data['show']['clauses'])
        html = self._html(prop)
        self.assertIn('月結要三天', html)

    # ------------------------------------------------------------------
    def test_docx_export_keeps_chapters_and_tables(self):
        try:
            import docx
        except ImportError:                  # pragma: no cover
            self.skipTest('沒有 python-docx')
        prop = self._full()
        binding = self.env.ref(BINDING_XMLID)
        _html, trees = binding._build_report_html(prop)
        output = self.env['doc.output']._record_output(binding, prop, trees[prop.id])
        output.action_download_docx()
        attachment = self.env['ir.attachment'].search(
            [('res_model', '=', 'doc.output'), ('res_id', '=', output.id)])
        self.assertTrue(attachment)
        document = docx.Document(io.BytesIO(base64.b64decode(attachment.datas)))
        text = '\n'.join(p.text for p in document.paragraphs)
        for needle in ('文件客戶股份有限公司', 'v1.0', '一、需求核對與本版處理',
                       '預算規劃', '備註與免責', '效益評估'):
            self.assertIn(needle, text)
        data = prop.doc_report_values()
        expected_tables = sum(1 for key, show in (
            ('changes', 'changes'), ('drivers', 'drivers'), ('caps', 'solution'),
            ('matrix', 'matrix'), ('masters', 'migration'), ('scope', 'scope'),
            ('budget_units', 'budget_time'), ('budget_lines', 'budget_sub'),
            ('phases', 'phases'), ('units', 'units'), ('payments', 'payments'),
            ('obligations', 'obligations'), ('clauses', 'clauses'), ('benefits', 'benefits'))
            if data['show'].get(show) and data.get(key))
        got = [' | '.join(c.text.strip()[:10] for c in t.rows[0].cells) for t in document.tables]
        self.assertEqual(len(document.tables), expected_tables,
                         'DOCX 的表：%s；資料顯示：%s' % (got, data['show']))
        # 資料列都在：矩陣表的列數＝表頭＋每個對應一列
        matrix = next(t for t in document.tables if t.rows[0].cells[1].text.strip() == '貴公司需求')
        self.assertEqual(len(matrix.rows), 1 + len(data['matrix']))


    # ------------------------------------------------------------------
    # 下載 Word
    # ------------------------------------------------------------------
    def _docx_of(self, prop):
        import docx
        action = prop.action_download_word()
        self.assertEqual(action['type'], 'ir.actions.act_url')
        attachment = self.env['ir.attachment'].search(
            [('res_model', '=', 'doc.output')], order='id desc', limit=1)
        return attachment, docx.Document(io.BytesIO(base64.b64decode(attachment.datas)))

    def test_download_word_names_the_file_and_leaves_an_audit_record(self):
        prop = self._full()
        prop.action_send()
        before = self.env['doc.output'].sudo().search_count([('res_id', '=', prop.id)])
        attachment, _document = self._docx_of(prop)
        self.assertEqual(attachment.name, '文件客戶股份有限公司_服務建議書_v1.0.docx',
                         '預設檔名含「/」（PRO/2026/0001），要換成可用的檔名')
        outputs = self.env['doc.output'].sudo().search([('res_id', '=', prop.id)])
        self.assertEqual(len(outputs), before + 1, '每次下載留一筆輸出紀錄')
        self.assertEqual(outputs.sorted('id')[-1].output_format, 'docx')

    def test_download_word_shades_the_table_headers_like_the_existing_specs(self):
        import docx
        from docx.oxml.ns import qn
        prop = self._full()
        _attachment, document = self._docx_of(prop)
        self.assertTrue(document.tables)
        for table in document.tables:
            head = table.rows[0].cells[0]._tc.tcPr.find(qn('w:shd'))
            self.assertIsNotNone(head, '表頭沒有底色')
            self.assertEqual(head.get(qn('w:fill')), docx_polish.HEADER_FILL)
        self.assertTrue(docx)

    def test_draft_can_be_exported_too(self):
        prop = self._full()
        self.assertEqual(prop.state, 'draft')
        _attachment, document = self._docx_of(prop)
        self.assertIn('文件客戶股份有限公司', '\n'.join(p.text for p in document.paragraphs))


@tagged('post_install', '-at_install')
class TestDocxPolish(TransactionCase):

    def _sample(self):
        import docx
        document = docx.Document()
        table = document.add_table(rows=4, cols=2)
        for r, row in enumerate(table.rows):
            for c, cell in enumerate(row.cells):
                cell.text = '表頭%d' % c if r == 0 else '資料%d-%d' % (r, c)
        document.add_paragraph('段落')
        buf = io.BytesIO()
        document.save(buf)
        return buf.getvalue()

    def test_header_is_filled_and_white_bold_rows_alternate(self):
        import docx
        from docx.oxml.ns import qn
        out = docx_polish.polish(self._sample())
        table = docx.Document(io.BytesIO(out)).tables[0]

        def fill(cell):
            shd = cell._tc.tcPr.find(qn('w:shd')) if cell._tc.tcPr is not None else None
            return shd.get(qn('w:fill')) if shd is not None else None

        self.assertEqual([fill(c) for c in table.rows[0].cells], [docx_polish.HEADER_FILL] * 2)
        self.assertIsNone(fill(table.rows[1].cells[0]))
        self.assertEqual(fill(table.rows[2].cells[0]), docx_polish.ZEBRA_FILL)
        self.assertIsNone(fill(table.rows[3].cells[0]))
        run = table.rows[0].cells[0].paragraphs[0].runs[0]
        self.assertTrue(run.font.bold)
        self.assertEqual(str(run.font.color.rgb), 'FFFFFF')

    def test_shading_is_inserted_in_schema_order(self):
        """☠️ Word 對 tcPr 子元素順序很挑：w:shd 必須在 w:vAlign 之前，順序錯了整份檔案打不開。"""
        import docx
        from docx.enum.table import WD_ALIGN_VERTICAL
        from docx.oxml.ns import qn
        document = docx.Document()
        table = document.add_table(rows=3, cols=1)
        for row in table.rows:
            row.cells[0].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
        buf = io.BytesIO()
        document.save(buf)
        out = docx.Document(io.BytesIO(docx_polish.polish(buf.getvalue())))
        for row in (out.tables[0].rows[0], out.tables[0].rows[2]):
            tags = [child.tag.split('}')[1] for child in row.cells[0]._tc.tcPr]
            self.assertIn('shd', tags)
            self.assertIn('vAlign', tags)
            self.assertLess(tags.index('shd'), tags.index('vAlign'), tags)
        # tblBorders 同理：要在 tblLook 之前
        tbl_tags = [c.tag.split('}')[1] for c in out.tables[0]._tbl.tblPr]
        self.assertIn('tblBorders', tbl_tags)
        if 'tblLook' in tbl_tags:
            self.assertLess(tbl_tags.index('tblBorders'), tbl_tags.index('tblLook'), tbl_tags)
        self.assertEqual(out.tables[0]._tbl.tblPr.find(qn('w:tblBorders')).find(
            qn('w:top')).get(qn('w:color')), docx_polish.BORDER_COLOR)

    def test_polish_changes_the_look_only_never_the_text(self):
        import docx
        before = docx.Document(io.BytesIO(self._sample()))
        after = docx.Document(io.BytesIO(docx_polish.polish(self._sample())))
        text = lambda d: [[c.text for c in r.cells] for t in d.tables for r in t.rows]  # noqa: E731
        self.assertEqual(text(before), text(after))
        self.assertEqual([p.text for p in before.paragraphs], [p.text for p in after.paragraphs])

    def test_polish_is_idempotent(self):
        import docx
        from docx.oxml.ns import qn
        once = docx_polish.polish(self._sample())
        twice = docx_polish.polish(once)
        table = docx.Document(io.BytesIO(twice)).tables[0]
        cell = table.rows[0].cells[0]
        self.assertEqual(len(cell._tc.tcPr.findall(qn('w:shd'))), 1, '不能疊出兩個底色')
