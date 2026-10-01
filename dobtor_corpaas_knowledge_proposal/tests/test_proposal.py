# -*- coding: utf-8 -*-
import base64
import io
import json
from unittest.mock import patch

from odoo.exceptions import UserError
from odoo.tests import TransactionCase, tagged

from ..services import calc, presales_xlsx


@tagged('post_install', '-at_install')
class TestCalc(TransactionCase):

    def test_effort_formula(self):
        # days = base + ceil(qty / unit) × per_unit
        self.assertEqual(calc.effort_days(2, 0.5, 20, 45), 3.5)
        self.assertEqual(calc.effort_days(2, 0.5, 20, 40), 3.0)
        self.assertEqual(calc.effort_days(2, 0.5, 20, 0), 2.0)
        self.assertEqual(calc.effort_days(1, 1, 0, 3), 4.0)  # unit_size 0 視為 1

    def test_moving_average(self):
        self.assertEqual(calc.moving_average(2.0, [4.0], 5), 3.0)
        self.assertEqual(calc.moving_average(2.0, [4.0, 6.0, 8.0], 2), 7.0)
        self.assertEqual(calc.observed_base(5.0, 0.5, 10, 25), 3.5)
        self.assertEqual(calc.observed_base(1.0, 1.0, 1, 5), 0.0)

    def test_third_party_parse(self):
        amount, guessed = calc.parse_third_party_monthly('簡訊 NT$1,200/月\n金流手續費', 500)
        self.assertEqual(amount, 1700.0)
        self.assertTrue(guessed)

    def test_parse_color(self):
        self.assertEqual(presales_xlsx.parse_color('🟢 原生'), 'native')
        self.assertEqual(presales_xlsx.parse_color('🟠客製微調＋🔴客製開發'), 'tuning')
        self.assertEqual(presales_xlsx.parse_color('客製微調'), 'tuning')
        self.assertEqual(presales_xlsx.parse_color('🔵Dobtor現有 / 🟢原生'), 'dobtor')
        self.assertEqual(presales_xlsx.parse_color('⚪沿用現況'), 'as_is')
        self.assertFalse(presales_xlsx.parse_color('⚪待確認'), '待確認不能當成沿用現況')
        self.assertEqual(presales_xlsx.parse_color('🟣企業版'), 'custom')
        self.assertEqual(presales_xlsx.parse_color('🟢原生（待確認）'), 'native')
        self.assertFalse(presales_xlsx.parse_color(''))


def _workbook_bytes():
    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = '0.總覽'
    ws.append(['總覽'])
    ws.append(['#', '客戶提問／痛點'])
    ws.append([1, '不該被匯入'])
    header = ['#', '客戶提問／痛點', '現況說明／客戶補充', '解決方案', '現有 or 客製',
              '對應原生功能／模組', '報價重點／待確認']
    fin = wb.create_sheet('1.財務')
    fin.append(['財務部痛點'])
    fin.append(header)
    fin.append([1, '月結要三天', 'Excel 對帳', '自動對帳', '🟢原生', 'account', ''])
    fin.append([2, None, None, None, None, None, None])
    fin.append([3, '發票要客製格式', '', '報表調整', '🟠客製微調＋🔴客製開發', '', '需確認格式'])
    hr = wb.create_sheet('2.人資')
    hr.append(['人資'])
    hr.append(header)
    hr.append([1, '請假沒系統', '紙本', '', '⚪沿用現況', '', ''])
    skip = wb.create_sheet('待確認事項')
    skip.append(['x'])
    skip.append(header)
    skip.append([1, '也不該匯入'])
    # 總覽的變體名稱、延伸分析表（表頭不同）都要略過
    ov = wb.create_sheet('總覽(更新)')
    ov.append(['總覽'])
    ov.append(header)
    ov.append([1, '總覽變體不該匯入'])
    ext = wb.create_sheet('9.訂單引擎拆解')
    ext.append(['延伸'])
    ext.append(['步驟', '說明', 'Odoo 落點'])
    ext.append([1, '延伸表不該匯入', 'sale'])
    # 一個痛點對兩個解決方案：# 與痛點欄往下合併
    ops = wb.create_sheet('3.營運')
    ops.append(['營運'])
    ops.append(header)
    ops.append([1, '門市叫貨靠 LINE', '手寫', '採購申請', '🟢原生', 'purchase', ''])
    ops.append([None, None, '', '自動補貨', '🟠客製微調', 'stock', '要確認安全庫存'])
    ops.merge_cells('A3:A4')
    ops.merge_cells('B3:B4')
    ops.append([2, '總部要看各店毛利', '', '', '🟣企業版', '', ''])
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


@tagged('post_install', '-at_install')
class TestProposal(TransactionCase):

    @classmethod
    def setUpClass(cls):
        super().setUpClass()
        icp = cls.env['ir.config_parameter'].sudo()
        for key, value in {'day_rate': 10000, 'internal_day_cost': 4000,
                           'cost_per_ccu_monthly': 100, 'cost_per_gb_monthly': 10,
                           'cost_per_worker_monthly': 1000, 'worker_mb': 512,
                           'third_party_default_monthly': 500, 'ai_point_cost': 2,
                           'calibration_window': 5, 'hours_per_day': 8}.items():
            icp.set_param('corpaas_proposal.%s' % key, value)
        cls.partner = cls.env['res.partner'].create({'name': '測試客戶'})
        Cap = cls.env['corpaas.knowledge.capability']
        cls.cap = Cap.create({
            'name': '線上報名', 'code': 'event_reg', 'color': 'native',
            'pain': '報名要人工整理', 'outcome': '自動收費',
            'master_data_models': 'res.partner',
            'third_party_costs': '金流 300/月',
            'resource_profile': json.dumps({'extra_worker_mb': 512}),
            'ai_points_monthly': 10,
        })
        cls.cap2 = Cap.create({'name': '會員管理', 'code': 'member', 'color': 'dobtor'})
        Tpl = cls.env['corpaas.knowledge.effort_template']
        cls.tpl_mig = Tpl.create({'capability_id': cls.cap.id, 'activity': 'migration',
                                  'base_days': 2, 'driver': 'records_100',
                                  'per_unit_days': 0.5, 'unit_size': 10})
        cls.tpl_train = Tpl.create({'capability_id': cls.cap.id, 'activity': 'training',
                                    'base_days': 1, 'driver': 'sessions',
                                    'per_unit_days': 0.5, 'unit_size': 1})
        cls.tpl_pm = Tpl.create({'activity': 'pm', 'base_days': 1, 'driver': 'users',
                                 'per_unit_days': 1, 'unit_size': 20})

    def _proposal(self):
        prop = self.env['corpaas.knowledge.proposal'].create({
            'partner_id': self.partner.id, 'ccu': 2, 'user_count': 45,
            'training_sessions': 2})
        pain = self.env['corpaas.knowledge.pain'].create({
            'proposal_id': prop.id, 'description': '報名要人工整理'})
        pain2 = self.env['corpaas.knowledge.pain'].create({
            'proposal_id': prop.id, 'description': '要客製報表'})
        self.env['corpaas.knowledge.mapping'].create([
            {'pain_id': pain.id, 'capability_id': self.cap.id, 'color': 'native',
             'confirmed': True},
            {'pain_id': pain2.id, 'color': 'custom', 'custom_days_low': 3,
             'custom_days_high': 5, 'confirmed': True},
            {'pain_id': pain2.id, 'capability_id': self.cap2.id, 'color': 'dobtor'},
        ])
        return prop

    def test_estimate_formula(self):
        prop = self._proposal()
        prop.action_compute_estimate()
        # 主資料檢核表由能力宣告的模型補齊
        md = prop.master_data_ids
        self.assertEqual(md.mapped('model'), ['res.partner'])
        md.record_count = 250
        prop.estimate_line_ids.unlink()
        self.env['corpaas.knowledge.estimate_line'].create(prop._estimate_vals())
        lines = {(l.template_id, l.activity): l for l in prop.estimate_line_ids}
        mig = lines[(self.tpl_mig, 'migration')]
        self.assertEqual(mig.driver_qty, 2.5)          # 250 筆 ÷ 100
        self.assertEqual(mig.days, 2 + 1 * 0.5)        # ceil(2.5/10)=1
        self.assertEqual(lines[(self.tpl_train, 'training')].days, 2.0)
        self.assertEqual(lines[(self.tpl_pm, 'pm')].days, 1 + 3 * 1)   # ceil(45/20)=3
        custom = prop.estimate_line_ids.filtered('is_custom')
        self.assertEqual(custom.days, 4.0)
        self.assertEqual(custom.amount, 40000)
        # 未確認的對應（會員管理）不進估算
        self.assertNotIn(self.cap2, prop.estimate_line_ids.mapped('capability_id'))
        self.assertEqual(prop.implementation_days, 2.5 + 2 + 4)
        self.assertEqual(prop.implementation_amount, 85000)
        self.assertEqual(prop.custom_amount, 40000)
        self.assertIn('報名要人工整理', prop.html)

    def test_margin(self):
        prop = self._proposal()
        prop.action_compute_estimate()
        prop.master_data_ids.record_count = 0
        # 人力：(2 + 2 + 4 + 4) 天 × 4000
        self.assertEqual(prop.labor_cost, 12 * 4000)
        # 基礎設施：2 CCU × 100 ＋ 1 worker × 1000
        self.assertAlmostEqual(prop.infra_monthly_cost, 1200)
        self.assertAlmostEqual(prop.third_party_monthly, 300)
        self.assertAlmostEqual(prop.ai_monthly_cost, 20)
        internal = 48000 + 12 * (1200 + 300 + 20)
        self.assertAlmostEqual(prop.internal_cost, internal)
        self.assertAlmostEqual(prop.total, 120000)  # 無訂閱方案：只有人天
        self.assertAlmostEqual(prop.margin, 120000 - internal)
        self.assertAlmostEqual(prop.margin_rate, (120000 - internal) / 120000 * 100)
        self.assertTrue(prop.cost_is_estimate)
        self.assertIn('估算值', prop.cost_html)

    def test_send_freezes(self):
        prop = self._proposal()
        prop.action_compute_estimate()
        prop.action_send()
        self.assertEqual(prop.state, 'sent')
        snap = json.loads(prop.snapshot_json)
        self.assertEqual(snap['prices']['total'], prop.total)
        self.assertEqual(len(snap['lines']), len(prop.estimate_line_ids))
        self.assertTrue(snap['html'])
        with self.assertRaises(UserError):
            prop.ccu = 5
        with self.assertRaises(UserError):
            prop.pain_ids[:1].description = '改掉'
        with self.assertRaises(UserError):
            prop.estimate_line_ids[:1].days = 99
        with self.assertRaises(UserError):
            prop.write({'snapshot_json': '{}'})
        with self.assertRaises(UserError):
            prop.action_compute_estimate()
        # ★ 送出後金額只讀快照：工時範本被校正、設定改了都不動
        total, cost = prop.total, prop.internal_cost
        self.tpl_mig.calibrate(9.0, 0)
        self.env['ir.config_parameter'].sudo().set_param('corpaas_proposal.cost_per_ccu_monthly', 999)
        prop.estimate_line_ids.invalidate_recordset()
        prop.invalidate_recordset()
        self.env.add_to_compute(prop._fields['total'], prop)
        self.env.add_to_compute(prop._fields['internal_cost'], prop)
        self.assertEqual((prop.total, prop.internal_cost), (total, cost))
        prop.action_mark_won()
        self.assertEqual(prop.state, 'won')
        self.assertEqual((prop.total, prop.internal_cost), (total, cost))
        new = prop.copy()
        self.assertEqual(new.state, 'draft')
        self.assertEqual(len(new.pain_ids), 2)
        self.assertEqual(len(new.mapping_ids), 3)

    def test_xlsx_import(self):
        try:
            data = _workbook_bytes()
        except ImportError:
            self.skipTest('openpyxl 未安裝')
        rows = presales_xlsx.parse_workbook(presales_xlsx.load_workbook_bytes(data))
        self.assertEqual([r['description'] for r in rows],
                         ['月結要三天', '發票要客製格式', '請假沒系統',
                          '門市叫貨靠 LINE', '總部要看各店毛利'])
        self.assertEqual([r['department'] for r in rows], ['財務', '財務', '人資', '營運', '營運'])
        self.assertEqual([r['color'] for r in rows],
                         ['native', 'tuning', 'as_is', 'native', 'custom'])
        self.assertEqual(rows[1]['quote_note'], '需確認格式')
        merged = rows[3]
        self.assertEqual(merged['proposed_solution'], '採購申請\n自動補貨')
        self.assertEqual(merged['module_hint'], 'purchase\nstock')
        self.assertEqual(merged['quote_note'], '要確認安全庫存')
        prop = self.env['corpaas.knowledge.proposal'].create({'partner_id': self.partner.id})
        wiz = self.env['corpaas.knowledge.proposal.import'].create({
            'proposal_id': prop.id, 'file': base64.b64encode(data), 'filename': 'x.xlsx'})
        wiz.action_import()
        self.assertEqual(len(prop.pain_ids), 5)
        self.assertEqual(set(prop.pain_ids.mapped('source')), {'import'})
        wiz2 = self.env['corpaas.knowledge.proposal.import'].create({
            'proposal_id': prop.id, 'file': base64.b64encode(data), 'replace': True})
        wiz2.action_import()
        self.assertEqual(len(prop.pain_ids), 5)

    def test_feedback_moving_average(self):
        prop = self._proposal()
        prop.action_compute_estimate()
        prop.action_send()
        rows = [('[event_reg] 資料移轉 客戶資料', 24.0),   # 3 天
                ('event_reg training 第一場', 8.0),
                ('週會', 8.0, ['event_reg', '教育訓練']),      # 任務標籤整個當一個詞
                ('雜項會議', 5.0),                             # 沒有工項
                ('資料移轉 對帳', 8.0),                         # 沒有能力 code → 不列入
                ('event_registration 資料移轉', 8.0),          # code 是子字串，不算
                ('通用 專案管理 週會', 8.0)]                    # 明確標通用 → 通用工項
        actuals = prop._actuals_from_rows(rows)
        self.assertEqual(actuals, {(self.cap.id, 'migration'): 3.0,
                                   (self.cap.id, 'training'): 2.0,
                                   (False, 'pm'): 1.0})
        done = prop._apply_actuals(actuals)
        self.assertEqual(done, 3)
        # migration：driver 0 → 反推基礎 3.0；(2 + 3) / 2
        self.assertEqual(self.tpl_mig.base_days, 2.5)
        self.assertEqual(self.tpl_mig.calibration_count, 1)
        hist = self.tpl_mig.calibration_ids
        self.assertEqual((hist.old_base_days, hist.new_base_days, hist.actual_days),
                         (2.0, 2.5, 3.0))
        # training：實際 2 天、2 場 → 扣掉 1.0，反推基礎 1.0；(1 + 1) / 2
        self.assertEqual(self.tpl_train.base_days, 1.0)
        # 通用 pm：實際 1 天、45 人 → 扣掉 3，反推基礎 0；(1 + 0) / 2
        self.assertEqual(self.tpl_pm.base_days, 0.5)
        # 同一建議書不重複校正
        self.assertEqual(prop._apply_actuals(actuals), 0)
        self.tpl_mig.calibrate(4.0, 0, window=5)
        self.assertEqual(self.tpl_mig.base_days, 3.0)  # (2 + 3 + 4) / 3

    def test_ai_match_mock(self):
        prop = self._proposal()
        tmpl = self.env['product.template'].create({'name': '測試方案', 'type': 'service'})
        pkg = self.env['infrastructure.solution.package'].create({'product_tmpl_id': tmpl.id})
        pkg.knowledge_capability_ids = [(6, 0, (self.cap | self.cap2).ids)]
        prop.package_id = pkg
        pain, pain2 = prop.pain_ids
        answer = {'mappings': [
            {'pain_id': pain.id, 'capability_id': self.cap.id, 'color': 'native',
             'custom_days': None, 'note': '用線上報名'},
            {'pain_id': pain2.id, 'capability_id': None, 'color': 'custom',
             'custom_days': [6, 2], 'note': '客製報表'},
            {'pain_id': 999999, 'capability_id': self.cap.id, 'color': 'native'},
            {'pain_id': pain2.id, 'capability_id': 999999, 'color': 'weird'},
        ]}
        Ai = type(self.env['corpaas.knowledge.ai'])
        calls = []

        def fake_enqueue(model_self, record, method_name, package, note=''):
            calls.append((record, method_name, package))
            return {'type': 'ir.actions.client', 'tag': 'display_notification', 'params': {}}
        with patch.object(Ai, 'enqueue', fake_enqueue), \
                patch.object(Ai, 'ask', return_value=answer) as ask:
            prop.action_ai_match()
            ask.assert_not_called()
            self.assertEqual(calls, [(prop, '_ai_match_run', pkg)])
            prop._ai_match_run()
        self.assertEqual(ask.call_args.args[0], 'proposal_match')
        ai = prop.mapping_ids.filtered(lambda m: m.source == 'ai')
        self.assertEqual(len(ai), 3)
        custom = ai.filtered(lambda m: m.color == 'custom' and m.custom_days_high)
        self.assertEqual((custom.custom_days_low, custom.custom_days_high), (2.0, 6.0))
        self.assertIn('2.0～6.0', custom.note)
        # 能力不在清單、顏色不合法 → 無能力、color=custom
        self.assertTrue(ai.filtered(lambda m: not m.capability_id and not m.custom_days_high))
        # 重跑：未確認的 AI 提議被取代，手動與已確認的保留
        ai[:1].confirmed = True
        with patch.object(Ai, 'ask', return_value={'mappings': []}):
            prop._ai_match_run()
        self.assertEqual(len(prop.mapping_ids.filtered(lambda m: m.source == 'ai')), 1)
        self.assertEqual(len(prop.mapping_ids.filtered(lambda m: m.source == 'manual')), 3)
        from odoo.addons.dobtor_corpaas_knowledge.services import hub_client
        with patch.object(Ai, 'ask', side_effect=hub_client.BudgetExceeded('滿了')):
            with self.assertRaises(UserError):
                prop._ai_match_run()

    def test_records_driver_falls_back_to_scenario_models(self):
        prop = self._proposal()
        self.env['corpaas.knowledge.proposal.master_data'].create(
            {'proposal_id': prop.id, 'model': 'res.partner', 'record_count': 300})
        Prop = type(prop)
        with patch.object(Prop, '_scenario_seed_models', lambda s: ['res.partner']):
            # 會員管理沒宣告主資料模型 → 用情境示範資料的模型
            self.assertEqual(prop._records_for(self.cap2), 300)
        with patch.object(Prop, '_scenario_seed_models', lambda s: []):
            self.assertEqual(prop._records_for(self.cap2), 0)

    def test_tuning_does_not_double_count_templates(self):
        prop = self._proposal()
        pain = prop.pain_ids[0]
        pain.mapping_ids.write({'color': 'tuning', 'custom_days_low': 2,
                                'custom_days_high': 4})
        vals = prop._estimate_vals()
        self.assertFalse([v for v in vals if v.get('activity') == 'tuning'],
                         '能力有工時範本時，設定微調不再加 AI 中位數')
        # 沒有能力（純設定微調）才用 AI 區間
        self.env['corpaas.knowledge.mapping'].create({
            'pain_id': pain.id, 'color': 'tuning', 'custom_days_low': 1,
            'custom_days_high': 3, 'confirmed': True})
        tuning = [v for v in prop._estimate_vals() if v.get('activity') == 'tuning']
        self.assertEqual([v['days'] for v in tuning], [2.0])

    def test_addons_all_modules_and_consistent_amounts(self):
        prop = self._proposal()
        Product = self.env['product.product']
        a = Product.create({'name': '加購 A', 'type': 'service', 'lst_price': 1000})
        b = Product.create({'name': '加購 B', 'type': 'service', 'lst_price': 2500})
        prop.mapping_ids.filtered('confirmed')[:1].addon_product_ids = [(6, 0, (a | b).ids)]
        self.assertEqual(prop._addon_products(), a | b)
        expected = prop._addon_first_year_for(a) + prop._addon_first_year_for(b)
        self.assertEqual(prop.addon_amount, expected)
        html = str(prop._render_html())
        for prod in (a, b):
            self.assertIn(prop._money(prop._addon_first_year_for(prod)), html)

    def _package_product(self):
        attr = self.env['product.attribute'].create({
            'name': 'KB 分支', 'create_variant': 'always',
            'value_ids': [(0, 0, {'name': '17.0'}), (0, 0, {'name': '18.0'})]})
        cycle_attr = self.env.ref('dobtor_corpaas_product.product_attribute_billing_cycle',
                                  raise_if_not_found=False)
        lines = [(0, 0, {'attribute_id': attr.id, 'value_ids': [(6, 0, attr.value_ids.ids)]})]
        if cycle_attr:
            vals = [self.env.ref('dobtor_corpaas_product.product_attr_cycle_value_%s' % c)
                    for c in ('monthly', 'yearly')]
            lines.append((0, 0, {'attribute_id': cycle_attr.id,
                                 'value_ids': [(6, 0, [v.id for v in vals])]}))
        tmpl = self.env['product.template'].create({
            'name': 'KB 版本方案', 'type': 'service', 'attribute_line_ids': lines})
        for v in tmpl.product_variant_ids:
            v.branch_name = v.product_template_attribute_value_ids.filtered(
                lambda x: x.attribute_id == attr).name
        return tmpl

    def test_package_variant_follows_package_and_branch(self):
        tmpl = self._package_product()
        v18 = tmpl.product_variant_ids.filtered(lambda v: v.branch_name == '18.0')
        target = v18.filtered(lambda v: v.paas_billing_cycle in (False, 'monthly'))[:1]
        pkg = self.env['infrastructure.solution.package'].create(
            {'product_tmpl_id': tmpl.id, 'product_id': target.id})
        prop = self.env['corpaas.knowledge.proposal'].create({
            'partner_id': self.partner.id, 'product_tmpl_id': tmpl.id, 'package_id': pkg.id,
            'billing_cycle': 'monthly'})
        self.assertEqual(prop._package_variant(), target, '取套件指定的變體，不是 template 第一個')
        yearly = v18.filtered(lambda v: v.paas_billing_cycle == 'yearly')
        if yearly:
            prop.billing_cycle = 'yearly'
            self.assertEqual(prop._package_variant(), yearly, '週期兄弟只在同分支裡找')

    def test_quotation_does_not_auto_provision(self):
        prop = self._proposal()
        tmpl = self.env['product.template'].create({'name': 'KB 報價方案', 'type': 'service'})
        prop.product_tmpl_id = tmpl
        prop.action_compute_estimate()
        prop.action_create_sale_order()
        order = prop.sale_order_id
        self.assertEqual(order.knowledge_proposal_id, prop)
        self.assertEqual(order.knowledge_provision_state, 'hold')
        self.assertFalse(order.is_to_create_paas, '不能預設開新平台')
        self.assertEqual(order._knowledge_package_line().product_id, tmpl.product_variant_id)
        # 擋住：回 True（已處理）＝ action_confirm 連疊加／換層級分支都跳過，並留言
        self.assertTrue(order._maybe_provision_by_tier())
        self.assertIn('確認開通', order.message_ids[:1].body)
        self.assertFalse(order.contract_id)
        # 既有客戶：模式不給預設
        Wiz = type(self.env['corpaas.knowledge.provision.wizard'])
        inst = self.env['infrastructure.instance']
        with patch.object(Wiz, '_knowledge_existing_instances', lambda s, p: inst.browse([1])):
            defaults = self.env['corpaas.knowledge.provision.wizard'].with_context(
                default_order_id=order.id).default_get(['mode', 'order_id', 'partner_id'])
        self.assertFalse(defaults.get('mode'))
        wiz = self.env['corpaas.knowledge.provision.wizard'].with_context(
            default_order_id=order.id).create({})
        self.assertEqual(wiz.mode, 'new')
        wiz.action_apply()
        self.assertEqual(order.knowledge_provision_state, 'released')
        self.assertTrue(order.is_to_create_paas)
        # 放行後交回 dobtor_corpaas_sale（這張單沒有自計費方案行 → 回 False）
        self.assertFalse(order._maybe_provision_by_tier())
        with self.assertRaises(UserError):
            order.action_knowledge_confirm_provision()

    def test_confirmed_order_only_releases_new(self):
        prop = self._proposal()
        tmpl = self.env['product.template'].create({'name': 'KB 報價方案 2', 'type': 'service'})
        prop.product_tmpl_id = tmpl
        prop.action_create_sale_order()
        order = prop.sale_order_id
        order.action_confirm()
        self.assertEqual(order.state, 'sale')
        self.assertEqual(order.knowledge_provision_state, 'hold')
        wiz = self.env['corpaas.knowledge.provision.wizard'].with_context(
            default_order_id=order.id).create({})
        wiz.mode = 'change_tier'
        with self.assertRaises(UserError):
            wiz.action_apply()
        wiz.mode = 'new'
        wiz.action_apply()
        self.assertEqual(order.knowledge_provision_state, 'released')

    def test_proposal_uses_live_pitch(self):
        if 'corpaas.knowledge.pitch' not in self.env:
            self.skipTest('沒有安裝行銷出口')
        prop = self._proposal()
        tmpl = self.env['product.template'].create({'name': 'KB 行銷方案', 'type': 'service'})
        pkg = self.env['infrastructure.solution.package'].create({'product_tmpl_id': tmpl.id})
        feature = self.env['corpaas.knowledge.feature'].create({
            'feature_key': 'kb_prop.action:reg', 'module': 'kb_prop', 'kind': 'action',
            'anchor': 'reg', 'name': '報名', 'package_ids': [(6, 0, pkg.ids)]})
        self.cap.feature_ids = [(6, 0, feature.ids)]
        prop.write({'product_tmpl_id': tmpl.id, 'package_id': pkg.id})
        pitch = self.env['corpaas.knowledge.pitch'].create({
            'capability_id': self.cap.id, 'product_tmpl_id': tmpl.id, 'package_id': pkg.id,
            'headline': '核准的標題', 'body_html': '<p>核准的內文</p>'})
        pitch._replace_claims([{'text': '可線上報名', 'features': [feature.feature_key]}])
        pitch.knowledge_propose('claim')
        pitch.action_approve()
        pitch.write({'headline': '草稿標題', 'body_html': '<p>草稿內文</p>'})
        pitch.action_submit()
        html = str(prop._render_html())
        self.assertIn('核准的標題', html)
        self.assertIn('核准的內文', html)
        self.assertIn('可線上報名', html)
        self.assertNotIn('草稿', html)

    def test_state_and_snapshot_internal_only(self):
        """狀態與送出快照只能由按鈕寫；帶 context 旗標（RPC 可偽造）也不行。"""
        prop = self._proposal()
        for ctx in ({}, {'proposal_freeze_bypass': True},
                    {'corpaas_proposal_internal': 'corpaas_proposal_internal'}):
            for vals in ({'state': 'won'}, {'snapshot_json': '{"prices": {}}'}):
                with self.assertRaises(UserError):
                    prop.with_context(**ctx).write(vals)
        with self.assertRaises(UserError):
            self.env['corpaas.knowledge.proposal'].create(
                {'partner_id': self.partner.id, 'state': 'sent'})
        self.assertEqual(prop.state, 'draft')
        prop.action_compute_estimate()
        prop.action_send()
        with self.assertRaises(UserError):
            prop.with_context(proposal_freeze_bypass=True).write({'state': 'lost'})
        prop.action_mark_lost()
        self.assertEqual(prop.state, 'lost')

    def _ccu_package(self, lo, hi):
        return self.env['product.template'].create({
            # sale_ok=False：不觸發 dobtor_corpaas_website 的上架開通目標檢查
            'name': 'KB CCU 方案', 'type': 'service', 'is_package': True, 'bill_by_ccu': True,
            'sale_ok': False,
            'ccu_tier_line_ids': [(0, 0, {'service_tier': 'shared', 'ccu_min': lo,
                                          'ccu_max': hi, 'ccu_unit_price': 100})]})

    def test_ccu_clamped_to_tier_bounds(self):
        tmpl = self._ccu_package(3, 10)
        prop = self._proposal()
        prop.write({'product_tmpl_id': tmpl.id, 'tier': 'shared', 'ccu': 1})
        self.assertEqual(prop._ccu_limits(), (3, 10))
        self.assertEqual(prop._effective_ccu(), 3)
        self.assertIn('CCU 3', prop.pricing_note, '報價以下限計價')
        prop.action_compute_estimate()
        prop.action_send()
        self.assertEqual(prop.ccu, 3)
        self.assertEqual(json.loads(prop.snapshot_json)['ccu'], 3)
        over = self._proposal()
        over.write({'product_tmpl_id': tmpl.id, 'tier': 'shared', 'ccu': 20})
        over.action_compute_estimate()
        with self.assertRaises(UserError):
            over.action_send()
        with self.assertRaises(UserError):
            over.action_create_sale_order()
        self.assertEqual(over.state, 'draft')

    def test_quotation_after_send_matches_snapshot(self):
        prop = self._proposal()
        tmpl = self.env['product.template'].create({'name': 'KB 快照方案', 'type': 'service'})
        addon = self.env['product.product'].create(
            {'name': '加購 C', 'type': 'service', 'lst_price': 1000})
        prop.product_tmpl_id = tmpl
        prop.mapping_ids.filtered('confirmed')[:1].addon_product_ids = [(6, 0, addon.ids)]
        prop.action_compute_estimate()
        prop.action_send()
        addon.lst_price = 2000
        with self.assertRaisesRegex(UserError, '複製新版'):
            prop.action_create_sale_order()
        self.assertFalse(prop.sale_order_id)
        addon.lst_price = 1000
        prop.action_create_sale_order()
        self.assertTrue(prop.sale_order_id)
