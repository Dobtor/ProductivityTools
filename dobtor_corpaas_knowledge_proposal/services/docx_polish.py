# -*- coding: utf-8 -*-
"""Word 後處理：補上 dobtor_doc_editor 的 DOCX 轉換不支援的表格外觀。

既有的報價規格書（預建 v6、以賽亞 v2.0、SKF 7.4）都是：表頭整列填主題色＋白色粗體字、
內容列交錯淡色底、表格有明確的淡灰框線。doc_editor 的轉換器只認 colspan／rowspan，
不吃儲存格底色與框線設定（實測確認），所以這一步在 DOCX 產出之後，用 python-docx 補上。

★ 只動外觀（儲存格底色、表頭字色、框線），不改任何文字或結構 —— 失敗（沒有 python-docx、
  檔案壞掉）時呼叫端應該退回未處理的檔案，而不是讓下載失敗。
☠️ 手動塞進去的 XML 元素**順序必須符合 OOXML 規格**：Word 對 `tcPr`／`tblPr` 的子元素順序
   很挑，順序錯了不是掉格式，而是整份檔案開啟時報「無法讀取內容」。本機沒有 Word 可驗，
   所以用 `_insert_before` 照規格順序插入，並由測試鎖住順序。
"""
import io

from .doc_template import THEME

HEADER_FILL = THEME.lstrip('#')
ZEBRA_FILL = 'F2F7FB'
BORDER_COLOR = 'BFBFBF'

#: 規格裡排在目標元素「後面」的兄弟（CT_TcPr／CT_TblPr 的序列）；新元素插在其中任何一個之前
_TCPR_AFTER_SHD = ('noWrap', 'tcMar', 'textDirection', 'tcFitText', 'vAlign', 'hideMark')
_TBLPR_AFTER_BORDERS = ('shd', 'tblLayout', 'tblCellMar', 'tblLook', 'tblCaption',
                        'tblDescription')


def _insert_before(parent, child, successors):
    """把 child 插到 parent 裡第一個 successors 之前；沒有就附在最後。"""
    from docx.oxml.ns import qn
    for name in successors:
        found = parent.find(qn('w:%s' % name))
        if found is not None:
            found.addprevious(child)
            return child
    parent.append(child)
    return child


def _shade(cell, fill):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn('w:shd')):
        tc_pr.remove(old)
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), fill)
    _insert_before(tc_pr, shd, _TCPR_AFTER_SHD)


def _borders(table, color):
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    tbl_pr = table._tbl.tblPr
    for old in tbl_pr.findall(qn('w:tblBorders')):
        tbl_pr.remove(old)
    borders = OxmlElement('w:tblBorders')
    # 規格順序：top, left(start), bottom, right(end), insideH, insideV
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:%s' % edge)
        el.set(qn('w:val'), 'single')
        el.set(qn('w:sz'), '4')
        el.set(qn('w:space'), '0')
        el.set(qn('w:color'), color)
        borders.append(el)
    _insert_before(tbl_pr, borders, _TBLPR_AFTER_BORDERS)


def polish(docx_bytes, header_fill=HEADER_FILL, zebra_fill=ZEBRA_FILL,
           border_color=BORDER_COLOR):
    """回處理後的 DOCX bytes。"""
    import docx
    from docx.shared import RGBColor
    document = docx.Document(io.BytesIO(docx_bytes))
    for table in document.tables:
        _borders(table, border_color)
        for index, row in enumerate(table.rows):
            if index == 0:
                for cell in row.cells:
                    _shade(cell, header_fill)
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.bold = True
                            run.font.color.rgb = RGBColor.from_string('FFFFFF')
            elif index % 2 == 0:
                for cell in row.cells:
                    _shade(cell, zebra_fill)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()
