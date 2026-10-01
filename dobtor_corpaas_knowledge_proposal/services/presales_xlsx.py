# -*- coding: utf-8 -*-
"""售前痛點 xlsx 解析（純函式；openpyxl workbook 進、dict 出）。

每張部門表：第 1 列標題、第 2 列表頭、第 3 列起資料，欄位
`# | 客戶提問／痛點 | 現況說明／客戶補充 | 解決方案 | 現有 or 客製 | 對應原生功能／模組 | 報價重點／待確認`。

★ 哪些表是部門表：看表頭，不看表名。部門表的第 2 列一定有「痛點／客戶提問」欄與
  「現有 or 客製」欄；總覽（含各種改名：「0.總覽」「總覽(更新)」…）、彙總表、待確認事項、
  解決方案總表、延伸分析表（功能拆解、版型對照…）表頭都不是這樣，一律略過。表名黑名單只是保險。
★ 合併儲存格：同一個痛點對到多個解決方案時，售前會把「#」「客戶提問」往下合併。
  ☠️ read_only 模式讀不到合併範圍（只有左上角有值，其他列看起來是空的，整列被當成沒有痛點而丟掉）。
  所以用一般模式載入，垂直合併的值往下補，同一個合併範圍的多列併成一個痛點（解決方案等欄位換行串起來）。

四色對照（售前 xlsx 的圖例 → 能力四色）：
  🟢 原生 → native；🔵 Dobtor 現有 → dobtor；🟠 客製微調 → tuning；🔴 客製開發 → custom；
  ⚪ 沿用現況 → as_is；
  ⚪ 待確認 → 不給顏色（False）：還沒判斷，不能當成「沿用現況」被排除在估算外，要人工或 AI 判斷；
  🟣 企業版（Odoo Enterprise 專屬功能）→ custom：CorPaaS 是 CE，要客製或替代方案才做得到。
組合標籤取最先出現的那一色。
"""
import io
import re

#: 表名保險（表頭判斷之外）：含這些字的表一律略過
SKIP_SHEET_WORDS = ('總覽', '彙總', '總表', '待確認事項', '延伸', '附錄', '說明', '圖例')

#: (關鍵字, 四色)；同位置以長的優先（「客製微調」不可被「客製」吃掉、「⚪待確認」不可被「⚪」吃掉）。
COLOR_KEYWORDS = [
    ('⚪待確認', False), ('⚪ 待確認', False), ('待確認', False),
    ('🟢', 'native'), ('🔵', 'dobtor'), ('🟠', 'tuning'), ('🔴', 'custom'), ('⚪', 'as_is'),
    ('🟣', 'custom'), ('企業版', 'custom'), ('Enterprise', 'custom'),
    ('原生', 'native'),
    ('Dobtor現有', 'dobtor'), ('Dobtor 現有', 'dobtor'), ('Dobtor', 'dobtor'),
    ('客製微調', 'tuning'), ('微調', 'tuning'),
    ('客製開發', 'custom'), ('客製', 'custom'),
    ('沿用現況', 'as_is'), ('沿用', 'as_is'),
]

#: 欄位 → 表頭關鍵字（依表頭找欄位；找不到就用固定位置）
HEADER_KEYS = [
    ('seq', ('#',), 0),
    ('description', ('痛點', '客戶提問'), 1),
    ('current_state', ('現況',), 2),
    ('proposed_solution', ('解決方案',), 3),
    ('color_text', ('現有', '客製'), 4),
    ('module_hint', ('對應', '原生功能', '模組'), 5),
    ('quote_note', ('報價', '待確認'), 6),
]

#: 同一痛點多列時，這些欄位換行串起來
JOIN_FIELDS = ('current_state', 'proposed_solution', 'color_text', 'module_hint', 'quote_note')

_DEPT = re.compile(r'^\s*\d+\s*[.．、\-_ ]\s*(.+?)\s*$')


def parse_color(text):
    """「現有 or 客製」欄 → 四色代碼；組合標籤取最先出現的那一色；待確認 → False。"""
    if not text:
        return False
    text = str(text)
    best = None
    for kw, color in COLOR_KEYWORDS:
        idx = text.find(kw)
        if idx < 0:
            continue
        key = (idx, -len(kw))
        if best is None or key < best[0]:
            best = (key, color)
    return best[1] if best else False


def _cell(row, idx):
    if idx is None or idx >= len(row):
        return ''
    v = row[idx]
    if v is None:
        return ''
    return str(v).strip()


def _find_column(texts, key, words, used):
    for i, t in enumerate(texts):
        if i in used or not t:
            continue
        if key == 'seq':
            if t in ('#', '＃', '序', '項次', 'No', 'No.'):
                return i
            continue
        if key == 'color_text':
            if '現有' in t or ('客製' in t and 'or' in t.lower()):
                return i
            continue
        if any(w in t for w in words):
            return i
    return None


def _columns(header):
    """依表頭定位欄位；回傳 (cols, 是否為部門表)。表頭沒對上的欄位用固定位置。"""
    cols, used, matched = {}, set(), set()
    texts = [str(h or '').strip() for h in header]
    for key, words, fallback in HEADER_KEYS:
        found = _find_column(texts, key, words, used)
        if found is not None:
            matched.add(key)
        elif fallback not in used:
            found = fallback
        if found is not None:
            used.add(found)
        cols[key] = found
    return cols, {'description', 'color_text'} <= matched


def _department(sheet_name):
    m = _DEPT.match(sheet_name or '')
    return m.group(1) if m else (sheet_name or '').strip()


def _skip_by_name(title):
    return any(w in (title or '') for w in SKIP_SHEET_WORDS)


def _grid(ws):
    """工作表 → 值的二維清單；垂直合併的值往下補。回傳 (grid, {(列, 欄): 合併起點列})。"""
    grid = [list(r) for r in ws.iter_rows(values_only=True)]
    anchors = {}
    for rng in getattr(getattr(ws, 'merged_cells', None), 'ranges', ()) or ():
        if rng.max_row <= rng.min_row:
            continue
        top = rng.min_row - 1
        for c in range(rng.min_col - 1, rng.max_col):
            value = grid[top][c] if top < len(grid) and c < len(grid[top]) else None
            for r in range(rng.min_row - 1, min(rng.max_row, len(grid))):
                row = grid[r]
                if c >= len(row):
                    row.extend([None] * (c + 1 - len(row)))
                row[c] = value
                anchors[(r, c)] = top
    return grid, anchors


def parse_workbook(wb):
    """回傳 [{'department','seq','description','current_state','proposed_solution',
    'color','color_text','module_hint','quote_note','sheet','row'}]。"""
    out = []
    for ws in wb.worksheets:
        if _skip_by_name(ws.title):
            continue
        if getattr(ws, 'reset_dimensions', None) and not getattr(ws, 'merged_cells', None):
            # ★ read_only 工作表：尺寸資訊常是錯的（只讀到第一列），要先重設
            ws.reset_dimensions()
        grid, anchors = _grid(ws)
        if len(grid) < 3:
            continue
        cols, is_dept = _columns(grid[1])
        if not is_dept:
            continue
        dept = _department(ws.title)
        desc_col = cols['description']
        by_anchor = {}
        for n, row in enumerate(grid[2:], start=2):
            desc = _cell(row, desc_col)
            if not desc:
                continue
            item = {
                'department': dept,
                'seq': _cell(row, cols['seq']),
                'description': desc,
                'current_state': _cell(row, cols['current_state']),
                'proposed_solution': _cell(row, cols['proposed_solution']),
                'color_text': _cell(row, cols['color_text']),
                'module_hint': _cell(row, cols['module_hint']),
                'quote_note': _cell(row, cols['quote_note']),
                'sheet': ws.title,
                'row': n + 1,
            }
            anchor = anchors.get((n, desc_col))
            prev = by_anchor.get(anchor) if anchor is not None else None
            if prev:
                # 同一個合併的痛點：其他欄位串起來（重複的不重複串）
                for key in JOIN_FIELDS:
                    if item[key] and item[key] not in prev[key].split('\n'):
                        prev[key] = '\n'.join(filter(None, [prev[key], item[key]]))
                continue
            if anchor is not None:
                by_anchor[anchor] = item
            out.append(item)
    for item in out:
        item['color'] = parse_color(item['color_text'])
    return out


def load_workbook_bytes(data):
    import openpyxl
    # ☠️ 不用 read_only：read_only 工作表沒有 merged_cells，合併的痛點會整列消失。
    return openpyxl.load_workbook(io.BytesIO(data), read_only=False, data_only=True)
