"""Tests for bulk import wizard (P3-1)。"""

import io
import base64
import zipfile

from importlib.util import find_spec

from odoo.exceptions import UserError
from odoo.tests.common import TransactionCase, tagged


def _make_minimal_docx_bytes():
    """產一個最小的 docx zip（只有必要 parts）。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        zf.writestr('[Content_Types].xml', '''<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
<Default Extension="xml" ContentType="application/xml"/>
<Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
<Override PartName="/word/document.xml" ContentType="application/vnd.openxmlformats-officedocument.wordprocessingml.document.main+xml"/>
</Types>''')
        zf.writestr('_rels/.rels', '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
<Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="word/document.xml"/>
</Relationships>''')
        zf.writestr('word/_rels/document.xml.rels', '''<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships"/>''')
        zf.writestr('word/document.xml', '''<?xml version="1.0" encoding="UTF-8"?>
<w:document xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main">
<w:body><w:p><w:r><w:t>Hello World</w:t></w:r></w:p></w:body>
</w:document>''')
    return buf.getvalue()


def _make_archive(entries):
    """打包多個 (name, bytes) → 一個外層 zip。"""
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as zf:
        for name, data in entries:
            zf.writestr(name, data)
    return buf.getvalue()


@tagged('post_install', '-at_install', 'dobtor_doc_import')
class TestBulkImport(TransactionCase):

    def setUp(self):
        super().setUp()
        self.Wizard = self.env['doc.bulk.import.wizard']
        self.Doc = self.env['doc.document']

    def test_wizard_creates_documents(self):
        """匯入含 2 個 docx 的 zip → 建立 2 筆文件。"""
        docx = _make_minimal_docx_bytes()
        archive = _make_archive([
            ('meeting1.docx', docx),
            ('meeting2.docx', docx),
        ])
        before = self.Doc.search_count([])

        wiz = self.Wizard.create({
            'archive_file': base64.b64encode(archive),
            'archive_filename': 'batch.zip',
            'target_company_id': self.env.company.id,
            'skip_failures': True,
        })
        wiz.action_run_import()

        self.assertEqual(wiz.state, 'done')
        self.assertEqual(wiz.created_count, 2)
        self.assertEqual(wiz.failed_count, 0)
        after = self.Doc.search_count([])
        self.assertEqual(after - before, 2)

    def test_wizard_skips_non_docx(self):
        """zip 內非 .docx 檔案應被略過，不影響 .docx 處理。"""
        archive = _make_archive([
            ('meeting.docx', _make_minimal_docx_bytes()),
            ('readme.txt', b'this should be skipped'),
            ('image.png', b'\x89PNG\r\n\x1a\n'),
        ])
        wiz = self.Wizard.create({
            'archive_file': base64.b64encode(archive),
            'archive_filename': 'mixed.zip',
            'target_company_id': self.env.company.id,
        })
        wiz.action_run_import()
        self.assertEqual(wiz.created_count, 1)
        self.assertIn('SKIP', wiz.log_text)

    def test_wizard_rejects_zip_bomb_attempt(self):
        """超過 5000 entry 的 archive 應被 zip_guard 擋。"""
        # 用最小 entry size 製造 5001 個假檔（會觸發 zip_guard 的 max_entries）
        entries = [(f'f{i}.txt', b'') for i in range(5001)]
        archive = _make_archive(entries)
        wiz = self.Wizard.create({
            'archive_file': base64.b64encode(archive),
            'archive_filename': 'too_many.zip',
            'target_company_id': self.env.company.id,
        })
        wiz.action_run_import()
        self.assertEqual(wiz.state, 'failed')
        self.assertIn('entry', (wiz.log_text or '').lower())

    def test_wizard_rejects_malformed_archive(self):
        wiz = self.Wizard.create({
            'archive_file': base64.b64encode(b'not a zip at all'),
            'archive_filename': 'broken.zip',
            'target_company_id': self.env.company.id,
        })
        wiz.action_run_import()
        self.assertEqual(wiz.state, 'failed')

    def test_wizard_assigns_target_company(self):
        """新建文件的 company_id 應 = wizard 設定的 target_company_id。"""
        before_ids = self.Doc.search([]).ids
        archive = _make_archive([('x.docx', _make_minimal_docx_bytes())])
        wiz = self.Wizard.create({
            'archive_file': base64.b64encode(archive),
            'archive_filename': 'one.zip',
            'target_company_id': self.env.company.id,
        })
        wiz.action_run_import()
        new_docs = self.Doc.search([('id', 'not in', before_ids)])
        self.assertEqual(len(new_docs), 1)
        self.assertEqual(new_docs.company_id, self.env.company)


@tagged('post_install', '-at_install', 'dobtor_doc_import', 'security')
class TestBulkImportCompanyBoundary(TransactionCase):
    """批次匯入不可以把文件建到使用者沒有的公司。

    ☠️ 2026-10-09 實測的真缺陷：`target_company_id` 沒有任何約束，而建立文件用
    `Doc.sudo().create(...)`。一個只屬於 A 公司的 doc manager 可以用 RPC 把
    target 設成他**連讀都讀不到**的 B 公司，文件就被建到 B 去了（實測 1 份）。

    和匯出紀錄那次（核心的 TestExportLogAccess）同一類：**`sudo()` 繞過
    record rule ＋ 使用者可控的公司欄位**。

    修法兩層，這組測試分別驗：
      1. 伺服端檢查 `target_company_id in env.companies` → 給可讀訊息
      2. create 去掉 `sudo()` → doc.document 的公司 rule 真的生效
    """

    def setUp(self):
        super().setUp()
        self.company_a = self.env.company
        self.company_b = self.env['res.company'].create({'name': '邊界測試公司B'})
        self.manager_a = self.env['res.users'].create({
            'name': 'A 公司的文件管理者',
            'login': 'bulk_boundary_mgr_a',
            'company_id': self.company_a.id,
            'company_ids': [(6, 0, [self.company_a.id])],
            'groups_id': [(6, 0, [
                self.env.ref('base.group_user').id,
                self.env.ref('dobtor_doc_editor.group_doc_manager').id,
            ])],
        })

    def _archive_with_one_docx(self):
        from docx import Document
        buf = io.BytesIO()
        doc = Document()
        doc.add_paragraph('邊界測試')
        doc.save(buf)
        zbuf = io.BytesIO()
        with zipfile.ZipFile(zbuf, 'w') as zf:
            zf.writestr('boundary.docx', buf.getvalue())
        return base64.b64encode(zbuf.getvalue())

    def test_cannot_import_into_a_company_the_user_does_not_have(self):
        """把 target 設成看不到的公司 → UserError，而且 B 公司一份都不該多。"""
        if find_spec('docx') is None:
            self.skipTest('需要 python-docx 才產得出測試用的 docx')
        before = self.env['doc.document'].sudo().search_count(
            [('company_id', '=', self.company_b.id)])
        wizard = self.env['doc.bulk.import.wizard'].with_user(self.manager_a).create({
            'archive_file': self._archive_with_one_docx(),
            'archive_filename': 'boundary.zip',
            'target_company_id': self.company_b.id,
        })
        with self.assertRaises(UserError):
            wizard.action_run_import()
        after = self.env['doc.document'].sudo().search_count(
            [('company_id', '=', self.company_b.id)])
        self.assertEqual(
            after, before,
            '文件被建到使用者沒有的公司去了——跨公司寫入成立（建了 %d 份）'
            % (after - before))

    def test_can_import_into_own_company(self):
        """自己的公司當然要能匯入——修正不可以把正常路徑擋掉。"""
        if find_spec('docx') is None:
            self.skipTest('需要 python-docx 才產得出測試用的 docx')
        wizard = self.env['doc.bulk.import.wizard'].with_user(self.manager_a).create({
            'archive_file': self._archive_with_one_docx(),
            'archive_filename': 'boundary.zip',
            'target_company_id': self.company_a.id,
        })
        wizard.action_run_import()
        self.assertEqual(wizard.created_count, 1,
                         '自己公司的匯入被擋住了：%s' % wizard.log_text)

    def test_target_company_field_has_a_domain(self):
        """UI 層也要擋——domain 只是第一層，但不能沒有。"""
        field = self.env['doc.bulk.import.wizard']._fields['target_company_id']
        self.assertTrue(
            field.domain,
            'target_company_id 沒有 domain，UI 會把所有公司都列出來')


@tagged('post_install', '-at_install', 'dobtor_doc_import')
class TestBulkImportCountsCorruptFilesAsFailed(TransactionCase):
    """壞檔要算進 failed_count，不可以建出一份「內容是錯誤訊息」的文件。

    ☠️ 2026-10-09 之前：`_docx_to_html_with_format()` 遇到「合法 zip 但不是
    docx」會回傳 `<p>（無法解析 DOCX：…）</p>`——也就是把失敗當內容。精靈於是
    把那份文件**建出來並計為成功**。改成 raise 之後，精靈逐檔的 except 會把它
    算進 failed_count 並寫進 log——精靈本來就是為這件事準備了那兩個欄位。
    """

    def setUp(self):
        super().setUp()
        if find_spec('docx') is None:
            self.skipTest('需要 python-docx')

    def _archive(self, entries):
        buf = io.BytesIO()
        with zipfile.ZipFile(buf, 'w') as zf:
            for name, blob in entries:
                zf.writestr(name, blob)
        return base64.b64encode(buf.getvalue())

    def _good_docx(self):
        from docx import Document
        b = io.BytesIO()
        d = Document()
        d.add_paragraph('好的文件')
        d.save(b)
        return b.getvalue()

    def _renamed_zip(self):
        b = io.BytesIO()
        with zipfile.ZipFile(b, 'w') as zf:
            zf.writestr('readme.txt', 'hello')
        return b.getvalue()

    def test_corrupt_entry_counts_as_failed_not_created(self):
        wizard = self.env['doc.bulk.import.wizard'].create({
            'archive_file': self._archive([
                ('good.docx', self._good_docx()),
                ('renamed.docx', self._renamed_zip()),
            ]),
            'archive_filename': 'mixed.zip',
            'skip_failures': True,
        })
        wizard.action_run_import()
        self.assertEqual(wizard.created_count, 1, '好的那份沒建出來')
        self.assertEqual(wizard.failed_count, 1,
                         '壞檔沒被算進 failed_count（log=%s）' % wizard.log_text)
        created = self.env['doc.document'].search(
            [('name', 'like', 'renamed')])
        self.assertFalse(
            created,
            '壞檔建出了文件——那份的內容會是錯誤訊息（id=%s）' % created.ids)
